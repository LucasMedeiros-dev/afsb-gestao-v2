"""
Cadastro dos perfis de cobrança a partir da "Lista para Lucas.xlsx" da AFSB.

Regras (combinadas com a AFSB):
- Aba "Lista": um perfil por associado, no CPF/CNPJ da coluna "BOLETO GERADO PARA".
- "BOLETO PARA CADA CNPJ": um perfil por CNPJ da aba "CNPJs BOLETO SEPARADO", com o
  valor daquele CNPJ (personalizado) e 20% de desconto até o vencimento.
- Vencimento dia 25 só para "VENC 25"; demais, dia 15.
- Cartão recorrente: perfil INATIVO (migração manual). Cartão 12x: perfil INATIVO
  (não gerar cobrança). PIX/boleto (Itaú ou Associatec): perfil ativo.
- Valor diferente da tabela por qtd. de lojas vira valor personalizado com motivo.

Não cadastra (vão para o relatório de pendências): associado fora do sistema,
pagador sem CPF/CNPJ e os números passados em --pular.
Nada é emitido no Asaas aqui. Idempotente por (associado, documento).
"""
import re
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone

from apps.cobranca.models import PerfilCobranca, TabelaValor
from apps.franquia.models import Franquia
from apps.usuario.models import Usuario

RE_DOC = re.compile(r'\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}|\d{3}\.\d{3}\.\d{3}-\d{2}')
digitos = lambda v: re.sub(r'\D', '', str(v or ''))  # noqa: E731


def dinheiro(v):
    return Decimal(str(v)).quantize(Decimal('0.01'))


class Command(BaseCommand):
    help = 'Cadastra os perfis de cobrança da planilha da AFSB (sem emitir nada no Asaas).'

    def add_arguments(self, parser):
        parser.add_argument('planilha')
        parser.add_argument('--pular', default='', help='Nºs de associado a não cadastrar (vírgula).')
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, planilha, pular, dry_run, **_):
        pular = {int(n) for n in pular.split(',') if n.strip()}
        wb = openpyxl.load_workbook(planilha, data_only=True)
        lista = [r for r in wb['Lista'].iter_rows(min_row=2, values_only=True) if any(c not in (None, '') for c in r)]
        separados = defaultdict(list)
        for r in wb['CNPJs BOLETO SEPARADO'].iter_rows(min_row=2, values_only=True):
            if any(c not in (None, '') for c in r):
                separados[int(r[0])].append(r)

        tabela = {t.qtd_lojas: (t.valor_cheio, t.valor_desconto) for t in TabelaValor.objects.filter(modalidade='mensal')}
        usuarios = {u.nro_associado: u for u in Usuario.objects.prefetch_related('franquias') if u.nro_associado}
        franquias = {digitos(f.cnpj): f for f in Franquia.objects.all()}
        inicio = timezone.localdate().replace(day=1)

        perfis, pulados = [], []
        for r in lista:
            nro, nome, _sit, lojas, meio, recebe, gerado, obs, cheio, desc = r[:10]
            nro, lojas, meio_u = int(nro), int(lojas), str(meio).upper()
            u = usuarios.get(nro)
            if nro in pular:
                pulados.append((nro, nome, 'aguardando confirmação da AFSB'))
                continue
            if not u:
                pulados.append((nro, nome, 'associado não cadastrado no sistema'))
                continue

            if 'RECORRENTE' in meio_u:
                forma, ativo, nota = 'cartao_recorrente', False, 'Cartão recorrente: inativo até migrar do sistema atual.'
            elif '12X' in meio_u:
                forma, ativo, nota = 'cartao_parcelado', False, f'Pago no cartão 12x ({meio.split("|")[-1].strip()}): não gerar cobrança.'
            else:
                forma, ativo, nota = 'pix_boleto', True, ''
            origem = 'Associatec' if 'ASSOCIATEC' in meio_u else 'Itaú' if 'ITAÚ' in meio_u else 'Cartão'
            obs_txt = ' | '.join(filter(None, [f'Origem: {origem}', f'Recebe por: {recebe}', str(obs or ''), nota]))
            base = {
                'usuario': u, 'forma_pagamento': forma, 'ativo': ativo,
                'dia_vencimento': 25 if 'VENC 25' in meio_u else 15,
                'qtd_parcelas': 12 if forma == 'cartao_parcelado' else None, 'observacao': obs_txt,
            }

            if str(gerado or '').strip().upper() == 'BOLETO PARA CADA CNPJ':
                for s in separados.get(nro, []):
                    valor = dinheiro(s[4])
                    doc = digitos(s[2])
                    perfis.append({
                        **base, 'tipo_documento': 'cnpj', 'documento': s[2], 'nome_pagador': str(s[1])[:100],
                        'qtd_lojas': 1, 'franquias': [franquias[doc]] if doc in franquias else [],
                        'personalizado': (valor, (valor * Decimal('0.8')).quantize(Decimal('0.01')),
                                          f'Boleto separado por CNPJ (cód. {s[3]}), valor conforme planilha da AFSB'),
                    })
                continue

            achado = RE_DOC.search(str(gerado or ''))
            if achado:
                doc = achado.group()
                tipo = 'cnpj' if len(digitos(doc)) == 14 else 'cpf'
            elif forma == 'pix_boleto':
                pulados.append((nro, nome, f'pagador sem CPF/CNPJ ("{gerado}")'))
                continue
            else:
                doc, tipo = u.cpf, 'cpf'  # Cartão: não emite boleto; documento só identifica o pagador.
            loja = franquias.get(digitos(doc))
            pagador = loja.razao_social if loja and tipo == 'cnpj' else (u.get_full_name() or nome)
            cheio, desc = dinheiro(cheio), dinheiro(desc)
            perfis.append({
                **base, 'tipo_documento': tipo, 'documento': doc, 'nome_pagador': pagador[:100],
                'qtd_lojas': lojas, 'franquias': list(u.franquias.all()),
                'personalizado': None if (cheio, desc) == tabela.get(lojas)
                else (cheio, desc, str(obs or 'Valor conforme planilha da AFSB')),
            })

        criados = atualizados = 0
        with transaction.atomic():
            for p in perfis:
                franq = p.pop('franquias')
                pers = p.pop('personalizado')
                p.update(
                    valor_personalizado_cheio=pers[0] if pers else None,
                    valor_personalizado_desconto=pers[1] if pers else None,
                    valor_personalizado_inicio=inicio if pers else None,
                    valor_personalizado_fim=None,
                    motivo_valor_personalizado=pers[2] if pers else None,
                )
                perfil, criado = PerfilCobranca.objects.update_or_create(
                    usuario=p.pop('usuario'), documento=p['documento'], defaults=p
                )
                perfil.franquias.set(franq)
                criados += criado
                atualizados += not criado
            if dry_run:
                transaction.set_rollback(True)

        for nro, nome, motivo in pulados:
            self.stdout.write(f'  pulado {nro} {nome}: {motivo}')
        self.stdout.write(self.style.SUCCESS(
            f"{'[dry-run] ' if dry_run else ''}{criados} perfis criados, {atualizados} atualizados, "
            f'{sum(1 for p in perfis if p["ativo"])} ativos; {len(pulados)} associados pulados.'
        ))
