"""
Carga inicial dos associados a partir da planilha da AFSB ("ASSOCIADOS dd.mm.aa.xlsx").

- Aba 1 (Relatório de Cadastros): um associado por linha + pares (nº da loja, CNPJ).
- Aba 2 (CNPJs): dados da empresa digitados à mão; só serve de reserva.
- Dados da empresa vêm do OpenCNPJ (cache em JSON para não repetir chamadas).

Associados com CPF ou e-mail repetido NÃO são importados: vão para o relatório
com os campos que divergem entre as linhas. Todo problema encontrado (CPF/CNPJ
inválido, WhatsApp fora do padrão, loja sem número, conflito de nº da loja...)
também vai para o relatório. Senha fica inutilizável: o primeiro acesso será
tratado depois.

Idempotente: atualiza por nº de associado (usuário) e por CNPJ (franquia).
"""
import json
import re
import time
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl.styles import Font, PatternFill

from apps.franquia.models import Franquia
from apps.usuario.models import Usuario

OPENCNPJ = 'https://api.opencnpj.org/{}'

# Colunas da aba 1 (0-based)
NRO, NOME, CPF, EMAIL, EMAIL2, WHATSAPP, TEL2, QTD_LOJAS, CONTRIB = range(9)
PRIMEIRA_LOJA, STATUS, FINANCEIRO, ADMISSAO, DESATIVACAO, MOTIVO = 9, 29, 30, 31, 32, 33
CAMPOS_ABA1 = {
    NRO: 'nº associado', NOME: 'nome', CPF: 'CPF', EMAIL: 'e-mail', EMAIL2: 'e-mail secundário',
    WHATSAPP: 'WhatsApp', TEL2: 'telefone secundário', QTD_LOJAS: 'qtd. lojas',
    CONTRIB: 'tipo de contribuição', STATUS: 'status', FINANCEIRO: 'situação financeira',
    ADMISSAO: 'data admissão', DESATIVACAO: 'data desativação', MOTIVO: 'motivo desativação',
}


def texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r'\s+', ' ', str(valor)).strip()


def digitos(valor):
    return re.sub(r'\D', '', texto(valor))


def cpf_valido(d):
    if len(d) != 11 or d == d[0] * 11:
        return False
    for n in (9, 10):
        resto = sum(int(d[i]) * (n + 1 - i) for i in range(n)) * 10 % 11 % 10
        if resto != int(d[n]):
            return False
    return True


def cnpj_valido(d):
    if len(d) != 14 or d == d[0] * 14:
        return False
    for n in (12, 13):
        pesos = list(range(n - 7, 1, -1)) + list(range(9, 1, -1))
        resto = sum(int(d[i]) * pesos[i] for i in range(n)) % 11
        if int(d[n]) != (0 if resto < 2 else 11 - resto):
            return False
    return True


def fmt_cpf(d):
    return f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'


def fmt_cnpj(d):
    return f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'


def fmt_cep(d):
    d = d.zfill(8)
    return f'{d[:5]}-{d[5:]}'


def data_br(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(texto(valor), '%d/%m/%Y').date()
    except ValueError:
        return None


class Relatorio:
    def __init__(self):
        self.pendencias = []
        self.duplicados = []

    def add(self, gravidade, tipo, associado='', nome='', campo='', valor='', detalhe=''):
        self.pendencias.append((gravidade, tipo, associado, nome, campo, valor, detalhe))

    def salvar(self, caminho):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Pendências'
        ws.append(['Gravidade', 'Tipo', 'Nº associado', 'Nome', 'Campo', 'Valor', 'Detalhe'])
        for linha in sorted(self.pendencias, key=lambda p: (p[0] != 'Não importado', p[1], str(p[2]))):
            ws.append(list(linha))

        ws2 = wb.create_sheet('Duplicados')
        ws2.append(['Grupo', 'Chave repetida'] + list(CAMPOS_ABA1.values()) + ['lojas', 'Campos divergentes'])
        destaque = PatternFill('solid', fgColor='FFF2CC')
        for grupo, chave, linhas, divergentes in self.duplicados:
            for linha in linhas:
                ws2.append([grupo, chave] + [texto(linha[i]) for i in CAMPOS_ABA1] + [
                    ', '.join(f'{l} ({c})' for l, c in pares_loja(linha)),
                    ', '.join(divergentes),
                ])
                for col, i in enumerate(CAMPOS_ABA1, start=3):
                    if CAMPOS_ABA1[i] in divergentes:
                        ws2.cell(ws2.max_row, col).fill = destaque

        for planilha in (ws, ws2):
            for celula in planilha[1]:
                celula.font = Font(bold=True)
            planilha.freeze_panes = 'A2'
            planilha.auto_filter.ref = planilha.dimensions
            for coluna in planilha.columns:
                largura = max(len(texto(c.value)) for c in coluna)
                planilha.column_dimensions[coluna[0].column_letter].width = min(max(largura + 2, 10), 60)
        wb.save(caminho)


def pares_loja(linha):
    pares = []
    for k in range(10):
        loja, cnpj = texto(linha[PRIMEIRA_LOJA + 2 * k]), texto(linha[PRIMEIRA_LOJA + 1 + 2 * k])
        if loja or cnpj:
            pares.append((loja, cnpj))
    return pares


class Command(BaseCommand):
    help = 'Importa associados (usuários) e lojas (franquias) da planilha da AFSB.'

    def add_arguments(self, parser):
        parser.add_argument('planilha')
        parser.add_argument('--relatorio', default='relatorio_importacao.xlsx')
        parser.add_argument('--cache', default='opencnpj_cache.json')
        parser.add_argument('--dry-run', action='store_true', help='Valida e gera o relatório sem gravar.')

    def handle(self, planilha, relatorio, cache, dry_run, **_):
        wb = openpyxl.load_workbook(planilha, data_only=True)
        aba1 = [r for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if any(texto(c) for c in r)]
        aba2 = [r[:12] for r in wb.worksheets[1].iter_rows(min_row=3, values_only=True) if any(texto(c) for c in r[:12])]
        rel = Relatorio()

        importaveis = self.separar_duplicados(aba1, rel)
        usuarios = self.preparar_usuarios(importaveis, rel)
        lojas = self.preparar_lojas(importaveis, aba2, rel, {u['nro_associado'] for u in usuarios})
        self.buscar_opencnpj(lojas, Path(cache), aba2, rel)

        with transaction.atomic():
            self.gravar(usuarios, lojas, rel)
            total_usuarios, total_franquias = Usuario.objects.count(), Franquia.objects.count()
            if dry_run:
                transaction.set_rollback(True)

        rel.salvar(relatorio)
        self.stdout.write(self.style.SUCCESS(
            f"{'[dry-run] ' if dry_run else ''}{total_usuarios} usuários, {total_franquias} franquias no banco. "
            f'{len(rel.pendencias)} pendências em {relatorio}'
        ))

    # ---------- usuários ----------

    def separar_duplicados(self, aba1, rel):
        """Agrupa linhas que compartilham CPF ou e-mail (união transitiva)."""
        pai = list(range(len(aba1)))

        def raiz(i):
            while pai[i] != i:
                pai[i] = pai[pai[i]]
                i = pai[i]
            return i

        por_chave = defaultdict(list)
        for i, r in enumerate(aba1):
            por_chave[('CPF', digitos(r[CPF]).zfill(11))].append(i)
            por_chave[('e-mail', texto(r[EMAIL]).lower())].append(i)
        chaves_grupo = defaultdict(set)
        for chave, idx in por_chave.items():
            if len(idx) > 1 and chave[1]:
                for j in idx[1:]:
                    pai[raiz(j)] = raiz(idx[0])
        for chave, idx in por_chave.items():
            if len(idx) > 1 and chave[1]:
                chaves_grupo[raiz(idx[0])].add(f'{chave[0]} {chave[1]}')

        grupos = defaultdict(list)
        for i in range(len(aba1)):
            grupos[raiz(i)].append(i)

        importaveis = []
        for n, (r, idx) in enumerate((r, idx) for r, idx in grupos.items() if len(idx) > 1):
            linhas = [aba1[i] for i in idx]
            divergentes = [
                nome for col, nome in CAMPOS_ABA1.items()
                if col != NRO and len({texto(l[col]).lower() for l in linhas}) > 1
            ]
            if len({tuple(pares_loja(l)) for l in linhas}) > 1:
                divergentes.append('lojas')
            rel.duplicados.append((n + 1, ' | '.join(sorted(chaves_grupo[r])), linhas, divergentes))
            for l in linhas:
                rel.add('Não importado', 'CPF/e-mail duplicado', texto(l[NRO]), texto(l[NOME]),
                        detalhe=f"Grupo {n + 1}; divergem: {', '.join(divergentes) or 'nada'}")
        for idx in grupos.values():
            if len(idx) == 1:
                importaveis.append(aba1[idx[0]])
        return importaveis

    def preparar_usuarios(self, linhas, rel):
        usuarios = []
        for r in linhas:
            nro, nome = int(texto(r[NRO])), texto(r[NOME])
            ctx = (nro, nome)

            cpf = digitos(r[CPF])
            if len(cpf) < 11 and cpf_valido(cpf.zfill(11)):
                rel.add('Corrigido', 'CPF sem zero à esquerda', *ctx, 'CPF', texto(r[CPF]), fmt_cpf(cpf.zfill(11)))
                cpf = cpf.zfill(11)
            if not cpf_valido(cpf):
                rel.add('Importado com erro', 'CPF inválido', *ctx, 'CPF', texto(r[CPF]))
            cpf_final = fmt_cpf(cpf) if len(cpf) == 11 else texto(r[CPF])

            wpp = digitos(r[WHATSAPP])
            wpp = wpp[2:] if len(wpp) in (12, 13) and wpp.startswith('55') else wpp
            if re.fullmatch(r'\d{2}[6-9]\d{7}', wpp):
                # Celular no formato antigo (8 dígitos): acrescenta o 9.
                rel.add('Corrigido', 'WhatsApp sem o 9', *ctx, 'WhatsApp', texto(r[WHATSAPP]),
                        f'+55 {wpp[:2]} 9{wpp[2:6]}-{wpp[6:]}')
                wpp = f'{wpp[:2]}9{wpp[2:]}'
            if not re.fullmatch(r'\d{2}9\d{8}', wpp):
                rel.add('Importado com erro', 'WhatsApp fora do padrão (DDD + 9 + 8 dígitos)', *ctx,
                        'WhatsApp', texto(r[WHATSAPP]), 'Provável celular sem o 9 ou fixo')

            status = {'ativo': Usuario.Status.ATIVO, 'desativado': Usuario.Status.INATIVO}.get(texto(r[STATUS]).lower())
            if status is None:
                rel.add('Importado com erro', 'Status desconhecido', *ctx, 'status', texto(r[STATUS]), 'Gravado como inativo')
                status = Usuario.Status.INATIVO
            desativado_em = texto(r[DESATIVACAO]) not in ('', '-')
            if status == Usuario.Status.ATIVO and (desativado_em or texto(r[MOTIVO])):
                rel.add('Divergência', 'Ativo com desativação preenchida', *ctx, 'status',
                        texto(r[STATUS]), f'{texto(r[DESATIVACAO])} {texto(r[MOTIVO])}'.strip())
            if status == Usuario.Status.INATIVO and not desativado_em:
                rel.add('Divergência', 'Desativado sem data de desativação', *ctx, 'status', texto(r[STATUS]))

            admissao = data_br(r[ADMISSAO])
            if admissao is None:
                rel.add('Importado com erro', 'Data de admissão inválida', *ctx, 'data admissão', texto(r[ADMISSAO]))

            qtd = texto(r[QTD_LOJAS])
            if qtd and qtd != str(len(pares_loja(r))):
                rel.add('Divergência', 'Qtd. de lojas ≠ lojas listadas', *ctx, 'qtd. lojas', qtd,
                        f'{len(pares_loja(r))} listada(s)')

            primeiro, _, resto = nome.partition(' ')
            usuarios.append({
                'nro_associado': nro,
                'email': texto(r[EMAIL]).lower(),
                'username': texto(r[EMAIL]).lower(),
                'first_name': primeiro[:150],
                'last_name': resto[:150],
                'cpf': cpf_final,
                'whatsapp': f'+55{wpp}' if wpp else '',
                'status': status,
                'data_admissao': admissao,
                'cargo': Usuario.Cargos.FRANQUEADO,
            })
        return usuarios

    # ---------- lojas ----------

    def preparar_lojas(self, linhas, aba2, rel, nros_importados):
        """CNPJ -> {nro_loja, associados}. Aba 1 manda; aba 2 completa o que faltar."""
        lojas = {}
        aba2_por_cnpj = {digitos(r[2]).zfill(14): r for r in aba2}

        def registrar(cnpj_txt, loja_txt, nro_assoc, nome, origem):
            ctx = (nro_assoc, nome)
            cnpj = digitos(cnpj_txt).zfill(14) if digitos(cnpj_txt) else ''
            if not cnpj:
                rel.add('Não importado', 'Loja sem CNPJ', *ctx, 'nº loja', loja_txt, origem)
                return
            if not cnpj_valido(cnpj):
                rel.add('Não importado', 'CNPJ inválido', *ctx, 'CNPJ', cnpj_txt, origem)
                return
            nro_loja = None
            if re.fullmatch(r'\d{1,3}(\.\d{3})*|\d+', loja_txt):
                nro_loja = int(loja_txt.replace('.', ''))
                if '.' in loja_txt:
                    rel.add('Corrigido', 'Nº da loja com ponto', *ctx, 'nº loja', loja_txt, str(nro_loja))
            elif loja_txt:
                rel.add('Não importado', 'Nº da loja não numérico', *ctx, 'nº loja', loja_txt, fmt_cnpj(cnpj))
                return
            loja = lojas.setdefault(cnpj, {'nro_loja': None, 'associados': set(), 'origem': origem})
            if nro_loja and loja['nro_loja'] and nro_loja != loja['nro_loja']:
                rel.add('Divergência', 'CNPJ com nº de loja diferente', *ctx, 'nº loja', str(nro_loja),
                        f"{fmt_cnpj(cnpj)} já registrado como loja {loja['nro_loja']} (mantido)")
            loja['nro_loja'] = loja['nro_loja'] or nro_loja
            if nro_assoc in nros_importados:
                loja['associados'].add(nro_assoc)

        for r in linhas:
            for loja_txt, cnpj_txt in pares_loja(r):
                registrar(cnpj_txt, loja_txt, int(texto(r[NRO])), texto(r[NOME]), 'aba 1')

        for cnpj, r in aba2_por_cnpj.items():
            nro_assoc = int(texto(r[0])) if texto(r[0]).isdigit() else None
            if cnpj in lojas:
                loja2 = digitos(r[3])
                if loja2 and lojas[cnpj]['nro_loja'] and int(loja2) != lojas[cnpj]['nro_loja']:
                    rel.add('Divergência', 'Nº da loja diferente entre abas', nro_assoc or '', '', 'nº loja',
                            texto(r[3]), f"{fmt_cnpj(cnpj)}: aba 1 = {lojas[cnpj]['nro_loja']} (mantido)")
                if nro_assoc in nros_importados:
                    lojas[cnpj]['associados'].add(nro_assoc)
            else:
                rel.add('Aviso', 'CNPJ só na aba 2', nro_assoc or '', '', 'CNPJ', texto(r[2]))
                registrar(texto(r[2]), texto(r[3]), nro_assoc, '', 'aba 2')
        return lojas

    def buscar_opencnpj(self, lojas, cache_path, aba2, rel):
        cache = json.loads(cache_path.read_text(encoding='utf-8')) if cache_path.exists() else {}
        aba2_por_cnpj = {digitos(r[2]).zfill(14): r for r in aba2}
        faltando = [c for c in lojas if c not in cache]
        for n, cnpj in enumerate(faltando, 1):
            cache[cnpj] = self.consultar(cnpj)
            if n % 50 == 0:
                self.stdout.write(f'OpenCNPJ: {n}/{len(faltando)}')
                cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding='utf-8')
        cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding='utf-8')

        for cnpj, loja in lojas.items():
            api = cache.get(cnpj)
            assoc = ', '.join(map(str, sorted(loja['associados'])))
            if api:
                if api.get('situacao_cadastral') != 'Ativa':
                    rel.add('Aviso', 'CNPJ não ativo na Receita', assoc, '', 'CNPJ', fmt_cnpj(cnpj),
                            api.get('situacao_cadastral', ''))
                logradouro = ' '.join(filter(None, [api.get('tipo_logradouro'), api.get('logradouro')]))
                loja['dados'] = {
                    'razao_social': api.get('razao_social', ''),
                    'nome_fantasia': api.get('nome_fantasia') or api.get('razao_social', ''),
                    'cep': fmt_cep(digitos(api.get('cep'))),
                    'rua': logradouro,
                    'numero': texto(api.get('numero')) or 'S/N',
                    'complemento': texto(api.get('complemento')) or None,
                    'bairro': api.get('bairro', ''),
                    'cidade': api.get('municipio', ''),
                    'estado': api.get('uf', ''),
                }
            elif cnpj in aba2_por_cnpj:
                r = aba2_por_cnpj[cnpj]
                rua, _, numero = texto(r[7]).rpartition(' - ')
                loja['dados'] = {
                    'razao_social': texto(r[4]),
                    'nome_fantasia': texto(r[5]) or texto(r[4]),
                    'cep': fmt_cep(digitos(r[6])),
                    'rua': rua or texto(r[7]),
                    'numero': numero if rua else 'S/N',
                    'complemento': texto(r[8]) or None,
                    'bairro': texto(r[9]),
                    'cidade': texto(r[10]),
                    'estado': texto(r[11]).upper(),
                }
                rel.add('Aviso', 'CNPJ fora do OpenCNPJ, usado aba 2', assoc, '', 'CNPJ', fmt_cnpj(cnpj))
            else:
                rel.add('Não importado', 'CNPJ fora do OpenCNPJ e da aba 2', assoc, '', 'CNPJ', fmt_cnpj(cnpj))

    def consultar(self, cnpj):
        for tentativa in range(5):
            try:
                req = urllib.request.Request(OPENCNPJ.format(cnpj), headers={'Accept': 'application/json', 'User-Agent': 'afsb-importador/1.0'})
                with urllib.request.urlopen(req, timeout=20) as resp:
                    return json.load(resp)
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    return None
                if e.code != 429 and e.code < 500:
                    raise
            except (urllib.error.URLError, TimeoutError):
                pass
            time.sleep(2 ** tentativa)
        raise RuntimeError(f'OpenCNPJ indisponível para {cnpj}')

    # ---------- gravação ----------

    def gravar(self, usuarios, lojas, rel):
        por_nro = {}
        for dados in usuarios:
            nro = dados.pop('nro_associado')
            usuario, criado = Usuario.objects.update_or_create(nro_associado=nro, defaults=dados)
            if criado:
                usuario.set_unusable_password()
                usuario.save(update_fields=['password'])
            por_nro[nro] = usuario

        usados = {}
        for cnpj, loja in sorted(lojas.items(), key=lambda kv: kv[1]['origem']):
            if not loja.get('dados'):
                continue
            assoc = ', '.join(map(str, sorted(loja['associados'])))
            if not loja['nro_loja']:
                rel.add('Não importado', 'Loja sem nº', assoc, '', 'CNPJ', fmt_cnpj(cnpj))
                continue
            if loja['nro_loja'] in usados:
                rel.add('Não importado', 'Nº da loja já usado por outro CNPJ', assoc, '', 'nº loja',
                        str(loja['nro_loja']), f"{fmt_cnpj(cnpj)} x {fmt_cnpj(usados[loja['nro_loja']])}")
                continue
            usados[loja['nro_loja']] = cnpj
            dados = {k: (v[:100] if isinstance(v, str) and k != 'numero' else v) for k, v in loja['dados'].items()}
            if len(dados['numero']) > 10:
                dados['complemento'] = ' '.join(filter(None, [dados['numero'], dados['complemento']]))[:100]
                dados['numero'] = dados['numero'][:10]
            franquia, _ = Franquia.objects.update_or_create(
                cnpj=fmt_cnpj(cnpj), defaults={'nro_da_loja': loja['nro_loja'], **dados}
            )
            franquia.usuarios.add(*(por_nro[n] for n in loja['associados'] if n in por_nro))
